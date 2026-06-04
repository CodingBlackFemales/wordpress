"""Extract structured text and images from python-pptx slides."""

from __future__ import annotations

import html
from dataclasses import dataclass, field
from pathlib import Path
from typing import Union

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.slide import Slide

from slides_to_learndash import blocks
from slides_to_learndash.rich_text import (
    classify_paragraph,
    paragraph_runs_to_html,
    paragraph_to_code_plain,
    plain_text_from_inner_html,
    strip_bold_tags_from_inner_html,
)
from slides_to_learndash.slide_geometry import build_row_column_plan, collect_body_shape_boxes

# Layout names that denote a section/divider slide (copyright footers appear here in many decks).
_SECTION_HEADER_MARKERS = frozenset({"SECTION_HEADER"})


def _layout_looks_like_section_header(layout_name: str) -> bool:
    n = layout_name.strip().upper().replace(" ", "_")
    return any(m in n for m in _SECTION_HEADER_MARKERS)


def _is_copyright_segment(kind: str, payload: str) -> bool:
    """Footer/legal lines that should not appear in lesson content."""
    if kind == "image":
        return False
    if kind == "code":
        plain = payload.strip().lower()
    else:
        plain = plain_text_from_inner_html(payload).strip().lower()
    if not plain:
        return False
    if "©" in plain or "(c)" in plain:
        return True
    if "all rights reserved" in plain:
        return True
    if "do not redistribute" in plain:
        return True
    if "copyright" in plain and ("reserved" in plain or "coding black females" in plain):
        return True
    return False


def _drop_copyright_segments(
    segments: list[tuple[str, str]], *, layout_name: str
) -> list[tuple[str, str]]:
    if not _layout_looks_like_section_header(layout_name):
        return segments
    return [s for s in segments if not _is_copyright_segment(s[0], s[1])]


@dataclass
class LinearBlock:
    """Single horizontal band (full width): ordered segments."""

    segments: list[tuple[str, str]]


@dataclass
class ColumnsBlock:
    """One wp:columns row."""

    widths: list[str]
    columns: list[list[tuple[str, str]]]


SlideContentBlock = Union[LinearBlock, ColumnsBlock]


@dataclass
class SlideExtract:
    index: int  # 1-based
    layout_name: str
    title: str
    content_blocks: list[SlideContentBlock] = field(default_factory=list)
    """Row-ordered blocks: linear segments or multi-column segments."""
    image_paths: list[str] = field(default_factory=list)

    @property
    def segments(self) -> list[tuple[str, str]]:
        """Flatten all segments (for topic title fallback and image paths)."""
        out: list[tuple[str, str]] = []
        for b in self.content_blocks:
            if isinstance(b, LinearBlock):
                out.extend(b.segments)
            else:
                for col in b.columns:
                    out.extend(col)
        return out


def _shape_rich_segments(shape) -> list[tuple[str, str]]:
    if not getattr(shape, "has_text_frame", False):
        return []
    out: list[tuple[str, str]] = []
    # Track blank paragraphs between code paragraphs: a blank acts as a block separator.
    blank_since_last_code = False
    for paragraph in shape.text_frame.paragraphs:
        kind = classify_paragraph(paragraph)
        if kind is None:
            blank_since_last_code = True
            continue
        if kind == "code":
            text = paragraph_to_code_plain(paragraph)
            # Merge into the previous code block only when it was truly adjacent
            # (no intervening blank paragraph).
            if out and out[-1][0] == "code" and not blank_since_last_code:
                out[-1] = ("code", out[-1][1] + "\n" + text)
            else:
                out.append(("code", text))
            blank_since_last_code = False
            continue
        blank_since_last_code = False
        if kind == "bullet":
            inner = paragraph_runs_to_html(paragraph)
            if not inner.strip():
                continue
            out.append(("bullet", inner))
        elif kind == "h3":
            inner = paragraph_runs_to_html(paragraph)
            if not inner.strip():
                continue
            out.append(("h3", inner))
        elif kind == "h4":
            inner = paragraph_runs_to_html(paragraph)
            if not inner.strip():
                continue
            out.append(("h4", inner))
        else:
            inner = paragraph_runs_to_html(paragraph)
            if not inner.strip():
                continue
            out.append(("p", inner))
    return out


def _segments_for_ordered_shapes(
    ordered_shapes: list,
    *,
    media_dir: Path | None,
    slide_index: int,
    image_counter: list[int],
    image_blob_cache: dict[bytes, str] | None = None,
    media_url_prefix: str = "media",
) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for shape in ordered_shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and media_dir is not None:
            rel = extract_picture(
                shape,
                media_dir=media_dir,
                slide_index=slide_index,
                image_counter=image_counter,
                image_blob_cache=image_blob_cache,
                media_url_prefix=media_url_prefix,
            )
            if rel:
                out.append(("image", rel))
            continue
        out.extend(_shape_rich_segments(shape))
    return out


def _slide_title_text(slide: Slide) -> str:
    if slide.shapes.title is not None:
        t = (slide.shapes.title.text or "").strip()
        if t:
            return t
    return ""


def extract_picture(
    shape,
    *,
    media_dir: Path,
    slide_index: int,
    image_counter: list[int],
    image_blob_cache: dict[bytes, str] | None = None,
    media_url_prefix: str = "media",
) -> str | None:
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return None
    image = shape.image
    blob = image.blob
    # Return the existing path if this exact image blob was already extracted.
    if image_blob_cache is not None and blob in image_blob_cache:
        return image_blob_cache[blob]
    ext = (image.ext or "png").lower()
    if ext not in ("png", "jpeg", "jpg", "gif", "webp", "tiff", "bmp"):
        ext = "png"
    image_counter[0] += 1
    n = image_counter[0]
    fname = f"slide_{slide_index:03d}_img_{n:02d}.{ext}"
    out = media_dir / fname
    out.write_bytes(blob)
    prefix = (media_url_prefix or "media").strip().strip("/")
    rel = f"{prefix}/{fname}"
    if image_blob_cache is not None:
        image_blob_cache[blob] = rel
    return rel


def _legacy_linear_blocks(
    slide: Slide,
    *,
    title_shape,
    media_dir: Path | None,
    slide_index: int,
    image_counter: list[int],
    image_blob_cache: dict[bytes, str] | None = None,
    media_url_prefix: str = "media",
) -> list[SlideContentBlock]:
    """Enumerate body shapes in file order (single column)."""
    segments: list[tuple[str, str]] = []
    for shape in slide.shapes:
        if title_shape is not None and shape == title_shape:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and media_dir is not None:
            rel = extract_picture(
                shape,
                media_dir=media_dir,
                slide_index=slide_index,
                image_counter=image_counter,
                image_blob_cache=image_blob_cache,
                media_url_prefix=media_url_prefix,
            )
            if rel:
                segments.append(("image", rel))
            continue
        segments.extend(_shape_rich_segments(shape))
    return [LinearBlock(segments)] if segments else []


def _flatten_all_segments(content_blocks: list[SlideContentBlock]) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for b in content_blocks:
        if isinstance(b, LinearBlock):
            out.extend(b.segments)
        else:
            for col in b.columns:
                out.extend(col)
    return out


def _apply_copyright_to_blocks(
    blocks_: list[SlideContentBlock], *, layout_name: str
) -> list[SlideContentBlock]:
    out: list[SlideContentBlock] = []
    for b in blocks_:
        if isinstance(b, LinearBlock):
            out.append(LinearBlock(_drop_copyright_segments(b.segments, layout_name=layout_name)))
        else:
            out.append(
                ColumnsBlock(
                    b.widths,
                    [_drop_copyright_segments(c, layout_name=layout_name) for c in b.columns],
                )
            )
    return out


def extract_slide(
    slide: Slide,
    slide_index: int,
    *,
    media_dir: Path | None,
    image_counter: list[int],
    image_blob_cache: dict[bytes, str] | None = None,
    media_url_prefix: str = "media",
) -> SlideExtract:
    layout_name = slide.slide_layout.name.strip()
    title = _slide_title_text(slide)
    title_shape = slide.shapes.title

    boxes = collect_body_shape_boxes(slide, title_shape=title_shape, media_dir=media_dir)

    if not boxes:
        content_blocks = _legacy_linear_blocks(
            slide,
            title_shape=title_shape,
            media_dir=media_dir,
            slide_index=slide_index,
            image_counter=image_counter,
            image_blob_cache=image_blob_cache,
            media_url_prefix=media_url_prefix,
        )
    else:
        plan = build_row_column_plan(boxes)
        content_blocks = []
        for cols, widths in plan:
            if len(cols) == 1:
                shapes = [sb.shape for sb in cols[0]]
                segs = _segments_for_ordered_shapes(
                    shapes,
                    media_dir=media_dir,
                    slide_index=slide_index,
                    image_counter=image_counter,
                    image_blob_cache=image_blob_cache,
                    media_url_prefix=media_url_prefix,
                )
                content_blocks.append(LinearBlock(segs))
            else:
                col_segments: list[list[tuple[str, str]]] = []
                for col in cols:
                    shapes = [sb.shape for sb in col]
                    col_segments.append(
                        _segments_for_ordered_shapes(
                            shapes,
                            media_dir=media_dir,
                            slide_index=slide_index,
                            image_counter=image_counter,
                            image_blob_cache=image_blob_cache,
                            media_url_prefix=media_url_prefix,
                        )
                    )
                content_blocks.append(ColumnsBlock(widths, col_segments))

    content_blocks = _apply_copyright_to_blocks(content_blocks, layout_name=layout_name)

    is_section_header = _layout_looks_like_section_header(layout_name)
    if is_section_header:
        # Section divider slides: only the slide title (header) is used elsewhere as a heading;
        # omit body text and all images (e.g. footer brand icons).
        content_blocks = []

    flat = _flatten_all_segments(content_blocks)
    if not flat and title:
        if not is_section_header:
            content_blocks = [LinearBlock([("p", html.escape(title))])]
            flat = _flatten_all_segments(content_blocks)

    rel_images = [s[1] for s in flat if s[0] == "image"]
    return SlideExtract(
        index=slide_index,
        layout_name=layout_name,
        title=title,
        content_blocks=content_blocks,
        image_paths=rel_images,
    )


def _segment_plain_for_dedupe(kind: str, payload: str) -> str:
    if kind == "image":
        return ""
    if kind == "code":
        return " ".join(payload.strip().split()).casefold()
    return plain_text_from_inner_html(payload).strip().casefold()


def dedupe_leading_title_duplicates(
    segments: list[tuple[str, str]],
    *,
    slide_title: str,
    include_slide_heading: bool,
) -> list[tuple[str, str]]:
    """Remove leading body segments that repeat the slide title (already emitted as a heading)."""
    if not include_slide_heading or not (slide_title or "").strip():
        return segments
    title_norm = slide_title.strip().casefold()
    out = list(segments)
    while out:
        k, pl = out[0]
        if k == "image":
            break
        plain = _segment_plain_for_dedupe(k, pl)
        if plain and plain == title_norm:
            out.pop(0)
            continue
        break
    return out


def segments_to_html_parts(segments: list[tuple[str, str]]) -> list[str]:
    """Convert segments to block HTML fragments (preserves order)."""
    parts: list[str] = []
    bullet_run: list[str] = []

    def flush_bullets() -> None:
        nonlocal bullet_run
        if bullet_run:
            parts.append(blocks.list_block_html(bullet_run))
            bullet_run = []

    for kind, payload in segments:
        if kind == "image":
            flush_bullets()
            parts.append(blocks.image_block(payload, alt=""))
        elif kind == "bullet":
            bullet_run.append(payload)
        elif kind == "code":
            flush_bullets()
            parts.append(blocks.code_block_plain(payload))
        elif kind == "h3":
            flush_bullets()
            parts.append(blocks.heading_block_html(strip_bold_tags_from_inner_html(payload), 3))
        elif kind == "h4":
            flush_bullets()
            parts.append(blocks.heading_block_html(strip_bold_tags_from_inner_html(payload), 4))
        elif kind == "p":
            flush_bullets()
            parts.append(blocks.paragraph_block_html(payload))
    flush_bullets()

    return [p for p in parts if p]


def slide_content_blocks_to_html(
    content_blocks: list[SlideContentBlock],
    *,
    slide_title: str,
    include_slide_heading: bool,
) -> list[str]:
    """Render row-ordered blocks to HTML fragments."""
    out: list[str] = []
    for b in content_blocks:
        if isinstance(b, LinearBlock):
            segs = dedupe_leading_title_duplicates(
                b.segments,
                slide_title=slide_title,
                include_slide_heading=include_slide_heading,
            )
            out.extend(segments_to_html_parts(segs))
        else:
            col_htmls: list[str] = []
            for col in b.columns:
                segs = dedupe_leading_title_duplicates(
                    col,
                    slide_title=slide_title,
                    include_slide_heading=include_slide_heading,
                )
                col_htmls.append(blocks.join_blocks(segments_to_html_parts(segs)))
            out.append(blocks.columns_block_html(col_htmls, b.widths))
    return [p for p in out if p]


def slide_to_post_html(
    sl: SlideExtract,
    *,
    include_slide_heading: bool = False,
    slide_heading_level: int = 3,
) -> str:
    """Build block HTML for one slide's content."""
    if slide_heading_level < 1:
        slide_heading_level = 1
    if slide_heading_level > 6:
        slide_heading_level = 6
    parts: list[str] = []
    if include_slide_heading:
        label = sl.title or f"Slide {sl.index}"
        parts.append(blocks.heading_block(label, level=slide_heading_level))
    parts.extend(
        slide_content_blocks_to_html(
            sl.content_blocks,
            slide_title=sl.title,
            include_slide_heading=include_slide_heading,
        )
    )
    return blocks.join_blocks(parts)


def load_presentation(path: Path) -> Presentation:
    return Presentation(str(path))

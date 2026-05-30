"""Geometry-based row/column layout for slide shapes (PPTX EMU coordinates)."""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Literal

from pptx.enum.shapes import MSO_SHAPE_TYPE

# Vertical overlap / min(height) above this ⇒ same row band.
_ROW_OVERLAP_MIN = 0.40
# Horizontal overlap / min(width) above this ⇒ same column (stacked shapes).
_COL_H_OVERLAP_MIN = 0.38
# Minimum horizontal gap (EMU) between column boxes to force separate columns when overlap is ambiguous.
_MIN_COLUMN_GAP_EMU = 36_000  # ~0.04 in at 914400 EMU/in
# Treat as side-by-side columns: horizontal overlap smaller than this (EMU) ⇒ not stacked in one column.
_MAX_X_OVERLAP_FOR_COLUMN_PAIR_EMU = 45_000  # ~0.05 in

SlideShapeKind = Literal["text", "picture"]


@dataclass
class ShapeBox:
    """Bounding box + reference to the underlying shape."""

    shape: Any
    left: int
    top: int
    width: int
    height: int
    kind: SlideShapeKind


def _bottom(sb: ShapeBox) -> int:
    return sb.top + sb.height


def _right(sb: ShapeBox) -> int:
    return sb.left + sb.width


def _vertical_overlap_ratio(a: ShapeBox, b: ShapeBox) -> float:
    """Intersection height / min(h1, h2)."""
    top_i = max(a.top, b.top)
    bot_i = min(_bottom(a), _bottom(b))
    inter = max(0, bot_i - top_i)
    h_min = min(a.height, b.height)
    if h_min <= 0:
        return 0.0
    return inter / h_min


def _x_interval_overlap_emu(a: ShapeBox, b: ShapeBox) -> int:
    """Width of horizontal intersection of two boxes (EMU)."""
    left_i = max(a.left, b.left)
    right_i = min(_right(a), _right(b))
    return max(0, right_i - left_i)


def _horizontal_overlap_ratio(a: ShapeBox, b: ShapeBox) -> float:
    """Intersection width / min(w1, w2)."""
    left_i = max(a.left, b.left)
    right_i = min(_right(a), _right(b))
    inter = max(0, right_i - left_i)
    w_min = min(a.width, b.width)
    if w_min <= 0:
        return 0.0
    return inter / w_min


def _side_by_side_column_pair(a: ShapeBox, b: ShapeBox) -> bool:
    """Two shapes that read as left/right columns: almost no horizontal overlap, but y-ranges intersect.

    Catches cases where vertical_overlap / min(height) is below _ROW_OVERLAP_MIN because one box is
    much taller than the other (short label beside a tall body text), which would otherwise split
    into separate rows and lose wp:columns.
    """
    if _x_interval_overlap_emu(a, b) > _MAX_X_OVERLAP_FOR_COLUMN_PAIR_EMU:
        return False
    top_i = max(a.top, b.top)
    bot_i = min(_bottom(a), _bottom(b))
    return bot_i > top_i


class _UnionFind:
    __slots__ = ("parent",)

    def __init__(self, n: int) -> None:
        self.parent = list(range(n))

    def find(self, x: int) -> int:
        p = self.parent
        while p[x] != x:
            p[x] = p[p[x]]
            x = p[x]
        return x

    def union(self, a: int, b: int) -> None:
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.parent[ra] = rb


def cluster_row_indices(boxes: list[ShapeBox]) -> list[list[int]]:
    """Group shape indices into rows using vertical overlap (union-find)."""
    n = len(boxes)
    if n == 0:
        return []
    if n == 1:
        return [[0]]
    uf = _UnionFind(n)
    for i in range(n):
        for j in range(i + 1, n):
            if _vertical_overlap_ratio(boxes[i], boxes[j]) >= _ROW_OVERLAP_MIN:
                uf.union(i, j)
    # Second pass: side-by-side column pairs with any vertical intersection (fixes short+tall column rows).
    for i in range(n):
        for j in range(i + 1, n):
            if _side_by_side_column_pair(boxes[i], boxes[j]):
                uf.union(i, j)
    buckets: dict[int, list[int]] = {}
    for i in range(n):
        r = uf.find(i)
        buckets.setdefault(r, []).append(i)
    rows = list(buckets.values())
    rows.sort(key=lambda idxs: min(boxes[i].top for i in idxs))
    for idxs in rows:
        idxs.sort(key=lambda i: boxes[i].left)
    return rows


def assign_columns_in_row(row_boxes: list[ShapeBox]) -> list[list[ShapeBox]]:
    """Split shapes in one row into side-by-side columns; merge stacked shapes into one column."""
    if len(row_boxes) <= 1:
        return [[row_boxes[0]]] if row_boxes else []
    ordered = sorted(row_boxes, key=lambda s: (s.left, s.top))
    columns: list[list[ShapeBox]] = []
    for sb in ordered:
        placed = False
        for col in columns:
            ref = col[0]
            hov = _horizontal_overlap_ratio(sb, ref)
            if hov >= _COL_H_OVERLAP_MIN:
                col.append(sb)
                col.sort(key=lambda s: s.top)
                placed = True
                break
            # Side-by-side: little x-overlap, but check gap from column bbox
            lo, hi = min(s.left for s in col), max(_right(s) for s in col)
            gap = sb.left - hi
            if gap >= _MIN_COLUMN_GAP_EMU and hov < 0.12:
                continue
        if not placed:
            columns.append([sb])
    columns.sort(key=lambda col: min(s.left for s in col))
    for col in columns:
        col.sort(key=lambda s: s.top)
    return columns


def column_width_fractions(columns: list[list[ShapeBox]]) -> list[float]:
    """Relative widths from column bounding boxes (sums to 1)."""
    if not columns:
        return []
    spans: list[int] = []
    for col in columns:
        lo = min(s.left for s in col)
        hi = max(_right(s) for s in col)
        spans.append(max(1, hi - lo))
    total = sum(spans)
    if total <= 0:
        n = len(columns)
        return [1.0 / n] * n
    return [s / total for s in spans]


_PRESETS_2 = (
    ("50%", "50%"),
    ("33.33%", "66.66%"),
    ("66.66%", "33.33%"),
)

_PRESETS_3 = (
    ("33.33%", "33.33%", "33.34%"),
    ("25%", "50%", "25%"),
)


def _dist(a: tuple[float, ...], b: tuple[float, ...]) -> float:
    return sum((x - y) ** 2 for x, y in zip(a, b, strict=True))


def widths_to_preset(fractions: list[float]) -> list[str]:
    """Map measured column fractions to the closest core Columns preset."""
    n = len(fractions)
    if n == 0:
        return []
    if n == 1:
        return ["100%"]
    if n == 2:
        f0, f1 = fractions[0], fractions[1]
        target = (f0, f1)
        best = _PRESETS_2[0]
        best_d = _dist(target, (0.5, 0.5))
        cand = ((0.3333, 0.6667), (0.6667, 0.3333))
        for i, t in enumerate(cand):
            d = _dist(target, t)
            if d < best_d:
                best_d = d
                best = _PRESETS_2[i + 1]
        return list(best)
    if n == 3:
        target = tuple(fractions)
        best = _PRESETS_3[0]
        best_d = _dist(target, (1 / 3, 1 / 3, 1 / 3))
        d2 = _dist(target, (0.25, 0.5, 0.25))
        if d2 < best_d:
            best = _PRESETS_3[1]
        return list(best)
    # 4+ columns: equal split (percent strings sum ~100)
    p = 100.0 / n
    return [f"{p:.2f}%" for _ in range(n)]


def build_row_column_plan(boxes: list[ShapeBox]) -> list[tuple[list[list[ShapeBox]], list[str]]]:
    """Return ordered rows; each row is (columns_of_shapeboxes, width_percent_strings)."""
    if not boxes:
        return []
    row_ixs = cluster_row_indices(boxes)
    plan: list[tuple[list[list[ShapeBox]], list[str]]] = []
    for ixs in row_ixs:
        row_boxes = [boxes[i] for i in ixs]
        cols = assign_columns_in_row(row_boxes)
        fr = column_width_fractions(cols)
        widths = widths_to_preset(fr)
        plan.append((cols, widths))
    return plan


def collect_body_shape_boxes(
    slide,
    *,
    title_shape,
    media_dir,
) -> list[ShapeBox]:
    """Non-title shapes that contribute text or exported pictures."""
    out: list[ShapeBox] = []
    for shape in slide.shapes:
        if title_shape is not None and shape == title_shape:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            if media_dir is not None:
                out.append(
                    ShapeBox(
                        shape=shape,
                        left=int(shape.left),
                        top=int(shape.top),
                        width=int(shape.width),
                        height=int(shape.height),
                        kind="picture",
                    )
                )
            continue
        if getattr(shape, "has_text_frame", False) and (shape.text_frame.text or "").strip():
            out.append(
                ShapeBox(
                    shape=shape,
                    left=int(shape.left),
                    top=int(shape.top),
                    width=int(shape.width),
                    height=int(shape.height),
                    kind="text",
                )
            )
    return out

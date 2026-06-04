"""CLI for slides-to-learndash."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Annotated, Optional

import typer

from slides_to_learndash.export_csv import write_csv
from slides_to_learndash.extract import load_presentation
from slides_to_learndash.pipeline import build_slides, run_lesson_only, run_lesson_with_topics

app = typer.Typer(add_completion=False, no_args_is_help=True)


def _list_layouts(input_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(input_path))
    typer.echo(f"File: {input_path}")
    typer.echo(f"Slides: {len(prs.slides)}\n")
    typer.echo(f"{'#':>4}  slide_layout.name")
    typer.echo("-" * 48)
    seen: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        name = slide.slide_layout.name
        typer.echo(f"{idx:4d}  {name}")
        if name not in seen:
            seen.append(name)
    typer.echo("\nDistinct layout names (first-seen order):")
    for n in seen:
        typer.echo(f"  {n}")


def _safe_dir_segment(name: str) -> str:
    """Filesystem-safe single path segment for per-export media folders."""
    s = name.strip()
    if not s:
        return "export"
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s)
    s = re.sub(r"\s+", "_", s).strip("._")
    return (s[:80] if s else "export") or "export"


def _default_media_dir(out_csv: Path) -> Path:
    """media/<sanitized_csv_stem>/ next to output CSV so each export has its own folder."""
    return out_csv.parent / "media" / _safe_dir_segment(out_csv.stem)


def _media_url_prefix(out_csv: Path, media_dir: Path) -> str:
    """Relative path from CSV directory to media dir, posix (matches src= in post HTML)."""
    try:
        rel = media_dir.resolve().relative_to(out_csv.parent.resolve())
        return str(rel).replace("\\", "/")
    except ValueError:
        return media_dir.name


def _resolve_media_dir(
    out: Path,
    media_dir: Optional[Path],
    skip_images: bool,
    multi_input: bool,
) -> tuple[Optional[Path], str]:
    """Return (mdir or None, media_prefix for build_slides)."""
    if skip_images:
        return None, "media"
    if media_dir is not None:
        base = media_dir.resolve()
        mdir = base / _safe_dir_segment(out.stem) if multi_input else base
        mdir.mkdir(parents=True, exist_ok=True)
        return mdir, _media_url_prefix(out, mdir)
    mdir = _default_media_dir(out)
    mdir.mkdir(parents=True, exist_ok=True)
    return mdir, _media_url_prefix(out, mdir)


def _export_lesson_only(
    input_path: Path,
    out: Path,
    course_id: str,
    section_heading_layout: Optional[str],
    media_dir: Optional[Path],
    skip_images: bool,
    lesson_title: Optional[str],
    slide_headings: bool,
    multi_input: bool,
) -> None:
    out = out.resolve()
    mdir, media_prefix = _resolve_media_dir(out, media_dir, skip_images, multi_input)

    prs = load_presentation(input_path)
    slides = build_slides(
        prs,
        media_dir=mdir,
        skip_images=skip_images,
        media_url_prefix=media_prefix,
    )

    rows = run_lesson_only(
        slides,
        course_id=course_id,
        lesson_title=lesson_title,
        include_slide_headings=slide_headings,
        section_heading_layout_regex=section_heading_layout,
    )
    write_csv(out, rows)
    typer.secho(f"Wrote {out} ({len(rows)} row(s)).", fg=typer.colors.GREEN)
    if mdir:
        typer.secho(f"Media: {mdir}")


def _export_lesson_with_topics(
    input_path: Path,
    out: Path,
    course_id: str,
    heading_layout_regex: str,
    media_dir: Optional[Path],
    skip_images: bool,
    lesson_title: Optional[str],
    slide_headings: bool,
    multi_input: bool,
) -> None:
    out = out.resolve()
    mdir, media_prefix = _resolve_media_dir(out, media_dir, skip_images, multi_input)

    prs = load_presentation(input_path)
    slides = build_slides(
        prs,
        media_dir=mdir,
        skip_images=skip_images,
        media_url_prefix=media_prefix,
    )

    lesson_rows, topic_rows = run_lesson_with_topics(
        slides,
        course_id=course_id,
        lesson_title=lesson_title,
        heading_layout_regex=heading_layout_regex,
        include_slide_headings=slide_headings,
    )
    lesson_path = out.parent / f"{out.stem}_lesson.csv"
    topics_path = out.parent / f"{out.stem}_topics.csv"
    write_csv(lesson_path, lesson_rows)
    typer.secho(f"Wrote {lesson_path} ({len(lesson_rows)} row(s)).", fg=typer.colors.GREEN)
    write_csv(topics_path, topic_rows)
    typer.secho(f"Wrote {topics_path} ({len(topic_rows)} row(s)).", fg=typer.colors.GREEN)
    typer.secho(
        "Import the lesson CSV with LearnDash Bulk → Content type Lessons, then fill lesson_id in the "
        "topics CSV from the new lesson post ID and import topics with Content type Topics.",
        fg=typer.colors.YELLOW,
    )
    if mdir:
        typer.secho(f"Media: {mdir}")


@app.command()
def main(
    ctx: typer.Context,
    pptx_files: Annotated[
        list[Path],
        typer.Argument(
            metavar="PPTX",
            help="One or more .pptx files (optional if --input is set instead)",
            exists=True,
            dir_okay=False,
            readable=True,
        ),
    ] = [],
    input_path: Optional[Path] = typer.Option(
        None,
        "--input",
        "-i",
        exists=True,
        dir_okay=False,
        readable=True,
        help="Path to a .pptx file (alternative to positional PPTX paths; do not combine both)",
    ),
    list_layouts: bool = typer.Option(
        False,
        "--list-layouts",
        help="Print each slide index and slide_layout.name; provide exactly one .pptx",
    ),
    mode: str = typer.Option(
        "lesson-only",
        "--mode",
        "-m",
        help="lesson-only (default): one merged lesson row; lesson-with-topics: lesson + topic rows",
    ),
    course_id: Optional[str] = typer.Option(
        None,
        "--course-id",
        help="LearnDash course_id column in the CSV (optional; leave unset for an empty value)",
    ),
    out: Optional[Path] = typer.Option(
        None,
        "--out",
        "-o",
        help="Output CSV path (.csv); default: same path as input with .csv extension. "
        "Not allowed with multiple inputs. For lesson-with-topics, writes {stem}_lesson.csv and {stem}_topics.csv",
    ),
    heading_layout: Optional[str] = typer.Option(
        None,
        "--heading-layout",
        help="Regex (Python) matched with re.fullmatch against slide_layout.name for topic boundaries",
    ),
    section_heading_layout: Optional[str] = typer.Option(
        None,
        "--section-heading-layout",
        help="Lesson-only: regex (re.fullmatch) for section/divider slides — slide title as H2; all others H3. "
        "Example: SECTION_HEADER",
    ),
    media_dir: Optional[Path] = typer.Option(
        None,
        "--media-dir",
        help="Directory for extracted images (default: <out parent>/media/<out_stem>/)",
    ),
    skip_images: bool = typer.Option(
        False,
        "--skip-images",
        help="Do not extract images from the PPTX",
    ),
    lesson_title: Optional[str] = typer.Option(
        None,
        "--lesson-title",
        help="Override lesson post_title (default: first slide title or text)",
    ),
    slide_headings: bool = typer.Option(
        True,
        "--slide-headings/--no-slide-headings",
        help="Before each slide's body, emit the slide title: H2 when the layout matches "
        "--section-heading-layout (lesson-only) or the topic heading layout (lesson-with-topics), else H3 (default: on)",
    ),
) -> None:
    """PPTX → LearnDash bulk-import CSV."""
    positional = list(pptx_files)
    from_option = [input_path] if input_path is not None else []

    if positional and from_option:
        typer.secho(
            "Error: use either positional .pptx file(s) or --input, not both.",
            fg=typer.colors.RED,
            err=True,
        )
        raise typer.Exit(code=2)

    inputs = positional if positional else from_option

    if list_layouts:
        if len(inputs) == 0:
            typer.secho(
                "Error: provide exactly one .pptx (positional or --input) with --list-layouts.",
                fg=typer.colors.RED,
                err=True,
            )
            raise typer.Exit(code=2)
        if len(inputs) > 1:
            typer.secho(
                "Error: --list-layouts accepts only one .pptx file.",
                fg=typer.colors.RED,
                err=True,
            )
            raise typer.Exit(code=2)
        p = inputs[0]
        if p.suffix.lower() != ".pptx":
            typer.secho("Error: file must be a .pptx.", fg=typer.colors.RED, err=True)
            raise typer.Exit(code=2)
        _list_layouts(p)
        return

    if len(inputs) == 0:
        typer.echo(ctx.get_help())
        raise typer.Exit(code=2)

    if len(inputs) > 1 and out is not None:
        typer.secho(
            "Error: --out is not supported with multiple input files (each output uses the input basename).",
            fg=typer.colors.RED,
            err=True,
        )
        raise typer.Exit(code=2)

    for p in inputs:
        if p.suffix.lower() != ".pptx":
            typer.secho(f"Error: not a .pptx file: {p}", fg=typer.colors.RED, err=True)
            raise typer.Exit(code=2)

    if mode not in ("lesson-only", "lesson-with-topics"):
        typer.secho("Error: --mode must be lesson-only or lesson-with-topics.", fg=typer.colors.RED, err=True)
        raise typer.Exit(code=2)

    if mode == "lesson-with-topics" and not (heading_layout and heading_layout.strip()):
        typer.secho("Error: --heading-layout is required for lesson-with-topics.", fg=typer.colors.RED, err=True)
        raise typer.Exit(code=2)

    cid = "" if course_id is None else str(course_id).strip()
    multi_input = len(inputs) > 1
    topic_heading_regex = heading_layout.strip() if heading_layout else ""

    for input_path in inputs:
        if out is not None and len(inputs) == 1:
            out_path = out.resolve()
        else:
            out_path = input_path.with_suffix(".csv")

        if mode == "lesson-only":
            _export_lesson_only(
                input_path=input_path,
                out=out_path,
                course_id=cid,
                section_heading_layout=section_heading_layout,
                media_dir=media_dir,
                skip_images=skip_images,
                lesson_title=lesson_title,
                slide_headings=slide_headings,
                multi_input=multi_input,
            )
        else:
            _export_lesson_with_topics(
                input_path=input_path,
                out=out_path,
                course_id=cid,
                heading_layout_regex=topic_heading_regex,
                media_dir=media_dir,
                skip_images=skip_images,
                lesson_title=lesson_title,
                slide_headings=slide_headings,
                multi_input=multi_input,
            )


if __name__ == "__main__":
    app()
